import pyodbc
import pandas as pd
import os
from datetime import datetime, timedelta
import re
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

print("=" * 80)
print("📊 SISTEMA DE ANÁLISE DE VOLUMES - HET E HOM")
print("=" * 80)

# Solicita a data inicial e final ao usuário no formato 20250829
data_inicial = input("Digite a data inicial (AAAAMMDD): ")
data_final = input("Digite a data final (AAAAMMDD): ")

# Converte para o formato YYYY-MM-DD
def formata_data(data):
    if len(data) == 8 and data.isdigit():
        return f"{data[:4]}-{data[4:6]}-{data[6:]}"
    else:
        print("⚠️  Digite a data no formato correto: AAAAMMDD")
        exit()

data_inicial_fmt = formata_data(data_inicial)
data_final_fmt = formata_data(data_final)

# Converte para datetime para cálculos
data_inicial_dt = datetime.strptime(data_inicial_fmt, "%Y-%m-%d")
data_final_dt = datetime.strptime(data_final_fmt, "%Y-%m-%d")

# Calcula range amplo (±2 dias)
data_ondaitm_inicial = (data_inicial_dt - timedelta(days=2)).strftime('%Y-%m-%d')
data_ondaitm_final = (data_final_dt + timedelta(days=2)).strftime('%Y-%m-%d')

# Função para limpar dados
def limpar_dados(df):
    for col in df.select_dtypes(include=['object']).columns:
        df[col] = df[col].astype(str).apply(lambda x: re.sub(r'[\x00-\x1f\x7f-\x9f]', '', x) if pd.notna(x) else x)
        df[col] = df[col].str.strip()
    return df

# Conexão
conn_str = (
    "DRIVER={iSeries Access ODBC Driver};"
    "SYSTEM=FGE5006CDP;"
    "UID=CDP174176G;"
    "PWD=CDP1753;"
)
conn = pyodbc.connect(conn_str)

print(f"\n📅 Período: {data_inicial_fmt} a {data_final_fmt}")
print("\n" + "=" * 80)
print("🔵 PROCESSANDO DADOS DE HET")
print("=" * 80)

# ==========================================
# PARTE 1: PROCESSAMENTO HET
# ==========================================

print("\n⏳ Carregando dados da GELIVE (HET)...")
query_gelive_het = f"""
SELECT *
FROM FGE5006CDP.GELIVE
WHERE MAJCRE >= '{data_inicial.replace("-", "")}'
    AND MAJCRE <= '{data_final.replace("-", "")}'
    AND (HEUEXC >= 140000 OR HEUEXC <= 010000)
    AND (REFLIV LIKE '2414%' OR REFLIV LIKE '2401%')
    AND CODTLI = 'STD'
"""

df_gelive_het = pd.read_sql(query_gelive_het, conn)
df_gelive_het = limpar_dados(df_gelive_het)
print(f"✅ GELIVE (HET) carregada: {len(df_gelive_het)} linhas")

# Limpar REFLIV e extrair VIAGEM e PALETE
df_gelive_het['REFLIV'] = df_gelive_het['REFLIV'].astype(str).str.split('.').str[0].str.strip()
df_gelive_het['VIAGEM'] = df_gelive_het['REFLIV'].str[4:-2].str.extract(r'(\d+)')[0].fillna(0).astype(int)
df_gelive_het['PALETE'] = df_gelive_het['REFLIV'].str[-2:].str.extract(r'(\d+)')[0].fillna(0).astype(int)
df_gelive_het['CODTLI'] = df_gelive_het['CODTLI'].astype(str).str.strip().str.upper()

# Remove duplicatas
df_gelive_het_unique = df_gelive_het.drop_duplicates(subset=['VIAGEM'], keep='first')
viagens_het = df_gelive_het['VIAGEM'].unique().tolist()
print(f"✅ {len(viagens_het)} viagens HET únicas encontradas")

if len(viagens_het) == 0:
    print("⚠️  Nenhuma viagem HET encontrada.")
    tons_preparados_het = 0
    tons_pendentes_het = 0
    df_filtrado_het = pd.DataFrame()
else:
    print(f"\n⏳ Carregando ONDAITM para HET...")
    viagens_str_het = ','.join([str(v) for v in viagens_het])
    query_ondaitm_het = f"""
    SELECT NUMVAG, PALETE, PREP, VIAGEM, CODTLI, STATUS, DONE, DATA, RCAIXAS, QTDLIDA, PESO
    FROM DANALLCDP.ONDAITM
    WHERE VIAGEM IN ({viagens_str_het})
        AND DATA >= '{data_ondaitm_inicial} 00:00:00'
        AND DATA <= '{data_ondaitm_final} 23:59:59'
        AND (CODTLI IS NULL OR TRIM(CODTLI) = '' OR CODTLI NOT LIKE '%HOM%')
    """

    df_ondaitm_het = pd.read_sql(query_ondaitm_het, conn)
    print(f"✅ ONDAITM (HET) carregada: {len(df_ondaitm_het)} linhas")

    # Processar status de preparação HET
    df_filtrado_het = df_ondaitm_het.dropna(how='all')

    def verificar_status(row):
        try:
            rcaixas = float(str(row['RCAIXAS']).replace(',', '.'))
            qtdlida = float(str(row['QTDLIDA']).replace(',', '.'))
            if rcaixas == qtdlida:
                return 'PREPARADO'
            else:
                return 'INCOMPLETO'
        except Exception:
            return 'ERRO'

    df_filtrado_het['STATUS_PREPARACAO'] = df_filtrado_het.apply(verificar_status, axis=1)

    # Calcular totais HET
    tons_preparados_het = df_filtrado_het[df_filtrado_het['STATUS_PREPARACAO'] == 'PREPARADO']['PESO'].sum()
    tons_pendentes_het = df_filtrado_het[df_filtrado_het['STATUS_PREPARACAO'] == 'INCOMPLETO']['PESO'].sum()

    print(f"\n📊 RESUMO HET:")
    print(f"   ✅ PREPARADO: {tons_preparados_het / 1000:.4f} tons ({tons_preparados_het:,.2f} kg)")
    print(f"   ⏳ PENDENTE: {tons_pendentes_het / 1000:.4f} tons ({tons_pendentes_het:,.2f} kg)")

# ==========================================
# PARTE 2: PROCESSAMENTO HOM
# ==========================================

print("\n" + "=" * 80)
print("🟢 PROCESSANDO DADOS DE HOM")
print("=" * 80)

print("\n⏳ Carregando dados da GELIVE (HOM - apenas STD)...")
# Para HOM: busca apenas viagens STD (mesmas que HET)
query_gelive_hom = f"""
SELECT *
FROM FGE5006CDP.GELIVE
WHERE MAJCRE >= '{data_inicial.replace("-", "")}'
    AND MAJCRE <= '{data_final.replace("-", "")}'
    AND (HEUEXC >= 140000 OR HEUEXC <= 010000)
    AND (REFLIV LIKE '2414%' OR REFLIV LIKE '2401%')
    AND CODTLI = 'STD'
"""

df_gelive_hom = pd.read_sql(query_gelive_hom, conn)
df_gelive_hom = limpar_dados(df_gelive_hom)
print(f"✅ GELIVE (HOM) carregada: {len(df_gelive_hom)} linhas")

# Processar GELIVE HOM
df_gelive_hom['REFLIV'] = df_gelive_hom['REFLIV'].astype(str).str.split('.').str[0].str.strip()
df_gelive_hom['VIAGEM'] = df_gelive_hom['REFLIV'].str[4:-2].str.extract(r'(\d+)')[0].fillna(0).astype(int)
df_gelive_hom['PALETE'] = df_gelive_hom['REFLIV'].str[-2:].str.extract(r'(\d+)')[0].fillna(0).astype(int)
df_gelive_hom['CODTLI'] = df_gelive_hom['CODTLI'].astype(str).str.strip().str.upper()

df_gelive_hom_unique = df_gelive_hom.drop_duplicates(subset=['VIAGEM'], keep='first')
print(f"   Viagens únicas: {df_gelive_hom['VIAGEM'].nunique()}")
print(f"   Tipos de CODTLI: {sorted(df_gelive_hom['CODTLI'].unique().tolist())}")

print(f"\n⏳ Carregando GESUPE (HOM)...")
data_gesupe_inicial = int((data_inicial_dt - timedelta(days=2)).strftime('%Y%m%d'))
data_gesupe_final = int((data_final_dt + timedelta(days=2)).strftime('%Y%m%d'))

query_gesupe = f"""
SELECT NUMSUP, TYPSUP, REFLIV, CARDES, MAJDAT, MAJCRE, MAJHMS
FROM FGE5006CDP.GESUPE
WHERE MAJCRE >= {data_gesupe_inicial}
    AND MAJCRE <= {data_gesupe_final}
    AND TYPSUP = 2
    AND (REFLIV LIKE '2414%' OR REFLIV LIKE '2401%')
"""

df_gesupe = pd.read_sql(query_gesupe, conn)
df_gesupe = limpar_dados(df_gesupe)
print(f"✅ GESUPE carregada: {len(df_gesupe)} linhas")

# Processar GESUPE
df_gesupe_2414 = df_gesupe[df_gesupe['REFLIV'].astype(str).str.strip().str.startswith('2414')].copy()
df_gesupe_2401 = df_gesupe[df_gesupe['REFLIV'].astype(str).str.strip().str.startswith('2401')].copy()
df_gesupe_final = pd.concat([df_gesupe_2414, df_gesupe_2401], ignore_index=True)

df_gesupe_final['VIAGEM'] = df_gesupe_final['REFLIV'].astype(str).str.strip().str[4:-2]
df_gesupe_final['STATUS'] = df_gesupe_final['CARDES'].apply(
    lambda x: 'PREPARADO' if pd.notna(x) and str(x).strip() != '' else 'PENDENTE'
)

# Merge com GELIVE para trazer CODTLI
df_gelive_hom_clean = df_gelive_hom_unique.copy()
df_gelive_hom_clean['VIAGEM'] = df_gelive_hom_clean['VIAGEM'].astype(str).str.strip()

df_gesupe_final = pd.merge(
    df_gesupe_final,
    df_gelive_hom_clean[['VIAGEM', 'CODTLI']],
    on='VIAGEM',
    how='inner'
)

print(f"✅ Merge concluído: {len(df_gesupe_final)} registros")
print(f"   Tipos de CODTLI encontrados após merge: {sorted(df_gesupe_final['CODTLI'].unique().tolist())}")

df_gesupe_final['TIPO_VIAGEM'] = df_gesupe_final['CODTLI']

# Para HOM: filtrar apenas STD
df_gesupe_hom = df_gesupe_final[df_gesupe_final['TIPO_VIAGEM'] == 'STD'].copy()
print(f"✅ GESUPE HOM final (apenas STD): {len(df_gesupe_hom)} registros")
if len(df_gesupe_hom) > 0:
    print(f"   Tipos de viagem: {sorted(df_gesupe_hom['TIPO_VIAGEM'].unique().tolist())}")

if len(df_gesupe_hom) == 0:
    print("⚠️  Nenhum registro HOM encontrado.")
    peso_preparado_hom = 0
    peso_pendente_hom = 0
    df_merge_hom = pd.DataFrame()
else:
    print(f"\n⏳ Carregando ONDAITM para HOM...")
    query_ondaitm_hom = f"""
    SELECT NUMSUP, PESO, VIAGEM
    FROM DANALLCDP.ONDAITM
    WHERE DATA >= '{data_ondaitm_inicial} 00:00:00'
        AND DATA <= '{data_ondaitm_final} 23:59:59'
    """

    df_ondaitm_hom = pd.read_sql(query_ondaitm_hom, conn)
    df_ondaitm_peso = df_ondaitm_hom[['NUMSUP', 'PESO']].drop_duplicates(subset=['NUMSUP'])
    print(f"✅ ONDAITM (HOM) carregada: {len(df_ondaitm_hom)} linhas, {len(df_ondaitm_peso)} suportes únicos")

    # Merge para trazer pesos
    df_merge_hom = pd.merge(df_gesupe_hom, df_ondaitm_peso, on='NUMSUP', how='left')

    registros_com_peso = df_merge_hom['PESO'].notna().sum()
    registros_sem_peso = df_merge_hom['PESO'].isna().sum()

    print(f"\n🔗 Correspondência NUMSUP (HOM):")
    print(f"   ✅ Com peso: {registros_com_peso}")
    print(f"   ⚠️  Sem peso: {registros_sem_peso}")

    # Calcular totais HOM (apenas STD)
    df_com_peso_hom = df_merge_hom[df_merge_hom['PESO'].notna()]
    
    peso_preparado_hom = df_com_peso_hom[df_com_peso_hom['STATUS'] == 'PREPARADO']['PESO'].sum()
    peso_pendente_hom = df_com_peso_hom[df_com_peso_hom['STATUS'] == 'PENDENTE']['PESO'].sum()
    peso_total_hom = df_com_peso_hom['PESO'].sum()

    print(f"\n📊 RESUMO HOM (apenas STD via GESUPE):")
    print(f"   ✅ PREPARADO: {peso_preparado_hom / 1000:.4f} tons ({peso_preparado_hom:,.2f} kg)")
    print(f"   ⏳ PENDENTE: {peso_pendente_hom / 1000:.4f} tons ({peso_pendente_hom:,.2f} kg)")
    print(f"   📦 TOTAL: {peso_total_hom / 1000:.4f} tons ({peso_total_hom:,.2f} kg)")

# ==========================================
# PARTE 3: CONSOLIDAÇÃO E TOTAIS
# ==========================================

print("\n" + "=" * 80)
print("📊 COMPARAÇÃO HET vs HOM (ambos apenas STD)")
print("=" * 80)
print("\n⚠️  IMPORTANTE: Mesmas viagens STD, metodologias diferentes!")
print("   - HET usa ONDAITM: Status baseado em RCAIXAS vs QTDLIDA")
print("   - HOM usa GESUPE: Status baseado em CARDES preenchido")
print("=" * 80)

print(f"\n🔵 HET (STD via ONDAITM):")
print(f"   ✅ Preparado: {tons_preparados_het / 1000:.4f} tons ({tons_preparados_het:,.2f} kg)")
print(f"   ⏳ Pendente: {tons_pendentes_het / 1000:.4f} tons ({tons_pendentes_het:,.2f} kg)")
print(f"   📦 Total: {(tons_preparados_het + tons_pendentes_het) / 1000:.4f} tons")

print(f"\n🟢 HOM (STD via GESUPE):")
print(f"   ✅ Preparado: {peso_preparado_hom / 1000:.4f} tons ({peso_preparado_hom:,.2f} kg)")
print(f"   ⏳ Pendente: {peso_pendente_hom / 1000:.4f} tons ({peso_pendente_hom:,.2f} kg)")
print(f"   📦 Total: {(peso_preparado_hom + peso_pendente_hom) / 1000:.4f} tons")

# Calcular totais consolidados
total_preparado = (tons_preparados_het + peso_preparado_hom) / 1000
total_pendente = (tons_pendentes_het + peso_pendente_hom) / 1000
total_geral = total_preparado + total_pendente

print(f"\n{'=' * 80}")
print(f"📈 TOTAL CONSOLIDADO (HET + HOM):")
print(f"{'=' * 80}")
print(f"   ✅ TOTAL PREPARADO: {total_preparado:.4f} tons ({tons_preparados_het + peso_preparado_hom:,.2f} kg)")
print(f"   ⏳ TOTAL PENDENTE: {total_pendente:.4f} tons ({tons_pendentes_het + peso_pendente_hom:,.2f} kg)")
print(f"   📦 TOTAL GERAL: {total_geral:.4f} tons ({tons_preparados_het + tons_pendentes_het + peso_preparado_hom + peso_pendente_hom:,.2f} kg)")
print(f"{'=' * 80}")

# ==========================================
# PARTE 4: EXPORTAÇÃO PARA EXCEL
# ==========================================

print("\n📄 Exportando dados para Excel...")
arquivo_saida = "Volume_HET_e_HOM_Consolidado.xlsx"

# Função para formatar o valor sem parênteses
def formatar_tons(valor):
    valor_formatado = valor / 1000
    return f"{valor_formatado:.4f}".replace('.', ',')

# Criar DataFrame de resumo consolidado
resumo_consolidado = pd.DataFrame({
    'FONTE': ['HET (ONDAITM)', 'HET (ONDAITM)', 'HOM (GESUPE)', 'HOM (GESUPE)', 'TOTAL HET + HOM', 'TOTAL HET + HOM'],
    'VIAGENS': ['STD', 'STD', 'STD', 'STD', 'STD', 'STD'],
    'STATUS': ['PREPARADO', 'PENDENTE', 'PREPARADO', 'PENDENTE', 'PREPARADO', 'PENDENTE'],
    'TONS': [
        formatar_tons(tons_preparados_het),
        formatar_tons(tons_pendentes_het),
        formatar_tons(peso_preparado_hom),
        formatar_tons(peso_pendente_hom),
        formatar_tons(tons_preparados_het + peso_preparado_hom),
        formatar_tons(tons_pendentes_het + peso_pendente_hom)
    ],
    'KG': [
        f"{tons_preparados_het:,.2f}",
        f"{tons_pendentes_het:,.2f}",
        f"{peso_preparado_hom:,.2f}",
        f"{peso_pendente_hom:,.2f}",
        f"{tons_preparados_het + peso_preparado_hom:,.2f}",
        f"{tons_pendentes_het + peso_pendente_hom:,.2f}"
    ]
})

with pd.ExcelWriter(arquivo_saida, engine='openpyxl') as writer:
    # Aba 1: Resumo Consolidado (primeira aba)
    resumo_consolidado.to_excel(writer, sheet_name='RESUMO_CONSOLIDADO', index=False)
    
    # Aba 2: Detalhes HET
    if not df_filtrado_het.empty:
        df_filtrado_het.to_excel(writer, sheet_name='HET_DETALHES', index=False)
    
    # Aba 3: Detalhes HOM
    if len(df_gesupe_hom) > 0:
        df_merge_hom.to_excel(writer, sheet_name='HOM_DETALHES', index=False)
    
    # Aba 4: GELIVE HET
    if not df_gelive_het.empty:
        df_gelive_het.to_excel(writer, sheet_name='GELIVE_HET', index=False)
    
    # Aba 5: GELIVE HOM
    if not df_gelive_hom.empty:
        df_gelive_hom.to_excel(writer, sheet_name='GELIVE_HOM', index=False)
    
    # Aba 6: GESUPE HOM
    if len(df_gesupe_hom) > 0:
        df_gesupe_hom.to_excel(writer, sheet_name='GESUPE_HOM', index=False)

# Adicionar filtros automáticos
wb = load_workbook(arquivo_saida)
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    if ws.max_row > 1:
        max_col_letter = get_column_letter(ws.max_column)
        ws.auto_filter.ref = f"A1:{max_col_letter}{ws.max_row}"

wb.save(arquivo_saida)
print(f"✅ Arquivo salvo em: {os.path.abspath(arquivo_saida)}")

# Fechar conexão
conn.close()

print("\n✅ Processamento concluído com sucesso!")
